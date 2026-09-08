import io
import os
from datetime import datetime
from django.conf import settings
from django.template.loader import render_to_string
from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, cm
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.piecharts import Pie
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import xhtml2pdf.pisa as pisa
from io import BytesIO

class ReportGenerator:
    """Generate various reports"""
    
    def __init__(self, report_type, data, template=None):
        self.report_type = report_type
        self.data = data
        self.template = template
        self.styles = getSampleStyleSheet()
        self._setup_styles()
    
    def _setup_styles(self):
        """Setup custom styles"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=24,
            alignment=TA_CENTER,
            spaceAfter=20,
            textColor=colors.HexColor('#1a5276')
        ))
        self.styles.add(ParagraphStyle(
            name='CustomHeading',
            parent=self.styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#2e86c1'),
            spaceAfter=10
        ))
        self.styles.add(ParagraphStyle(
            name='CustomBody',
            parent=self.styles['Normal'],
            fontSize=11,
            spaceAfter=6
        ))
    
    def generate_student_report_pdf(self, student, analysis, examination):
        """Generate student report as PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)
        story = []
        
        # Title
        story.append(Paragraph("STUDENT PERFORMANCE REPORT", self.styles['CustomTitle']))
        story.append(Spacer(1, 0.2*inch))
        
        # School Info
        story.append(Paragraph("EduGrade Analytics System", self.styles['CustomHeading']))
        story.append(Paragraph("School Performance Report", self.styles['CustomBody']))
        story.append(Spacer(1, 0.1*inch))
        
        # Student Information Table
        student_data = [
            ['Admission No:', student.admission_number],
            ['Name:', student.full_name],
            ['Grade/Form:', student.current_grade_level.name if student.current_grade_level else '-'],
            ['Stream:', student.current_stream.name if student.current_stream else '-'],
            ['Examination:', examination.name if examination else '-'],
            ['Date:', datetime.now().strftime('%Y-%m-%d')],
        ]
        student_table = Table(student_data, colWidths=[2*inch, 4*inch])
        student_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.black),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('PADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(student_table)
        story.append(Spacer(1, 0.2*inch))
        
        # Subject Performance
        if 'subject_performances' in self.data:
            story.append(Paragraph("Subject Performance", self.styles['CustomHeading']))
            subject_data = [['Subject', 'Score', 'Grade', 'Points']]
            for perf in self.data['subject_performances']:
                subject_data.append([
                    perf['subject_name'],
                    f"{perf['score']:.1f}%",
                    perf['grade'],
                    str(perf.get('points', '-'))
                ])
            
            subject_table = Table(subject_data, colWidths=[2*inch, 1.5*inch, 1.5*inch, 1.5*inch])
            subject_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2e86c1')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('PADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(subject_table)
            story.append(Spacer(1, 0.2*inch))
        
        # Overall Performance
        if 'overall' in self.data:
            overall = self.data['overall']
            story.append(Paragraph("Overall Performance", self.styles['CustomHeading']))
            overall_data = [
                ['Total Score', f"{overall.get('total_score', 0):.1f}"],
                ['Average', f"{overall.get('average_score', 0):.1f}%"],
                ['Grade', overall.get('grade', '-')],
                ['Position', str(overall.get('position', '-'))],
            ]
            overall_table = Table(overall_data, colWidths=[2*inch, 4*inch])
            overall_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('PADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(overall_table)
            story.append(Spacer(1, 0.2*inch))
        
        # Strengths and Weaknesses
        if 'strengths' in self.data and self.data['strengths']:
            story.append(Paragraph("Strengths", self.styles['CustomHeading']))
            for strength in self.data['strengths']:
                story.append(Paragraph(f"• {strength}", self.styles['CustomBody']))
            story.append(Spacer(1, 0.1*inch))
        
        if 'weaknesses' in self.data and self.data['weaknesses']:
            story.append(Paragraph("Areas for Improvement", self.styles['CustomHeading']))
            for weakness in self.data['weaknesses']:
                story.append(Paragraph(f"• {weakness}", self.styles['CustomBody']))
            story.append(Spacer(1, 0.1*inch))
        
        # Footer
        story.append(Spacer(1, 0.5*inch))
        story.append(Paragraph("Generated by EduGrade Analytics System", self.styles['Normal']))
        story.append(Paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}", self.styles['Normal']))
        
        doc.build(story)
        buffer.seek(0)
        return buffer
    
    def generate_class_report_pdf(self, class_data):
        """Generate class report as PDF"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)
        story = []
        
        # Title
        story.append(Paragraph("CLASS PERFORMANCE REPORT", self.styles['CustomTitle']))
        story.append(Spacer(1, 0.2*inch))
        
        # Class Info
        story.append(Paragraph(f"Class: {class_data.get('class_name', 'N/A')}", self.styles['CustomHeading']))
        story.append(Paragraph(f"Exam: {class_data.get('exam_name', 'N/A')}", self.styles['CustomBody']))
        story.append(Spacer(1, 0.1*inch))
        
        # Statistics
        stats = class_data.get('statistics', {})
        stats_data = [
            ['Total Students', str(stats.get('total_students', 0))],
            ['Average Score', f"{stats.get('average_score', 0):.1f}%"],
            ['Highest Score', f"{stats.get('highest_score', 0):.1f}%"],
            ['Lowest Score', f"{stats.get('lowest_score', 0):.1f}%"],
            ['Pass Rate', f"{stats.get('pass_rate', 0):.1f}%"],
        ]
        stats_table = Table(stats_data, colWidths=[3*inch, 3*inch])
        stats_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('PADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(stats_table)
        story.append(Spacer(1, 0.2*inch))
        
        # Top Students
        top_students = class_data.get('top_students', [])
        if top_students:
            story.append(Paragraph("Top Students", self.styles['CustomHeading']))
            top_data = [['Rank', 'Student', 'Score']]
            for i, student in enumerate(top_students[:10], 1):
                top_data.append([
                    str(i),
                    student.get('name', 'N/A'),
                    f"{student.get('score', 0):.1f}%"
                ])
            
            top_table = Table(top_data, colWidths=[1*inch, 3.5*inch, 2*inch])
            top_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2e86c1')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('PADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(top_table)
        
        # Footer
        story.append(Spacer(1, 0.5*inch))
        story.append(Paragraph("Generated by EduGrade Analytics System", self.styles['Normal']))
        story.append(Paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}", self.styles['Normal']))
        
        doc.build(story)
        buffer.seek(0)
        return buffer
    
    def generate_word_report(self, report_data, report_type='student'):
        """Generate Word document report"""
        doc = Document()
        
        # Add title
        title = doc.add_heading(report_data.get('title', 'Performance Report'), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add date
        doc.add_paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph()
        
        if report_type == 'student':
            self._add_student_section(doc, report_data)
        elif report_type == 'class':
            self._add_class_section(doc, report_data)
        else:
            self._add_general_section(doc, report_data)
        
        # Save to buffer
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer
    
    def _add_student_section(self, doc, data):
        """Add student section to Word document"""
        student = data.get('student', {})
        doc.add_heading('Student Information', level=1)
        doc.add_paragraph(f"Name: {student.get('name', 'N/A')}")
        doc.add_paragraph(f"Admission No: {student.get('admission', 'N/A')}")
        doc.add_paragraph(f"Grade: {student.get('grade', 'N/A')}")
        doc.add_paragraph(f"Stream: {student.get('stream', 'N/A')}")
        doc.add_paragraph()
        
        # Subject Performance
        subject_perf = data.get('subject_performances', [])
        if subject_perf:
            doc.add_heading('Subject Performance', level=1)
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Subject'
            hdr_cells[1].text = 'Score'
            hdr_cells[2].text = 'Grade'
            hdr_cells[3].text = 'Points'
            
            for perf in subject_perf:
                row_cells = table.add_row().cells
                row_cells[0].text = perf.get('subject_name', 'N/A')
                row_cells[1].text = f"{perf.get('score', 0):.1f}%"
                row_cells[2].text = perf.get('grade', '-')
                row_cells[3].text = str(perf.get('points', '-'))
            doc.add_paragraph()
        
        # Overall
        overall = data.get('overall', {})
        if overall:
            doc.add_heading('Overall Performance', level=1)
            doc.add_paragraph(f"Total Score: {overall.get('total_score', 0):.1f}")
            doc.add_paragraph(f"Average: {overall.get('average_score', 0):.1f}%")
            doc.add_paragraph(f"Grade: {overall.get('grade', '-')}")
            doc.add_paragraph()
        
        # Strengths and Weaknesses
        strengths = data.get('strengths', [])
        if strengths:
            doc.add_heading('Strengths', level=1)
            for strength in strengths:
                doc.add_paragraph(f"• {strength}")
            doc.add_paragraph()
        
        weaknesses = data.get('weaknesses', [])
        if weaknesses:
            doc.add_heading('Areas for Improvement', level=1)
            for weakness in weaknesses:
                doc.add_paragraph(f"• {weakness}")
    
    def _add_class_section(self, doc, data):
        """Add class section to Word document"""
        doc.add_heading('Class Information', level=1)
        doc.add_paragraph(f"Class: {data.get('class_name', 'N/A')}")
        doc.add_paragraph(f"Exam: {data.get('exam_name', 'N/A')}")
        doc.add_paragraph()
        
        stats = data.get('statistics', {})
        if stats:
            doc.add_heading('Statistics', level=1)
            doc.add_paragraph(f"Total Students: {stats.get('total_students', 0)}")
            doc.add_paragraph(f"Average Score: {stats.get('average_score', 0):.1f}%")
            doc.add_paragraph(f"Highest Score: {stats.get('highest_score', 0):.1f}%")
            doc.add_paragraph(f"Lowest Score: {stats.get('lowest_score', 0):.1f}%")
            doc.add_paragraph(f"Pass Rate: {stats.get('pass_rate', 0):.1f}%")
            doc.add_paragraph()
        
        top_students = data.get('top_students', [])
        if top_students:
            doc.add_heading('Top Students', level=1)
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Rank'
            hdr_cells[1].text = 'Student'
            hdr_cells[2].text = 'Score'
            
            for i, student in enumerate(top_students[:10], 1):
                row_cells = table.add_row().cells
                row_cells[0].text = str(i)
                row_cells[1].text = student.get('name', 'N/A')
                row_cells[2].text = f"{student.get('score', 0):.1f}%"
    
    def _add_general_section(self, doc, data):
        """Add general section to Word document"""
        doc.add_heading('Report Details', level=1)
        for key, value in data.items():
            if key not in ['title']:
                doc.add_paragraph(f"{key}: {value}")
